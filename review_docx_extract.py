from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(doc):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--end", type=int, default=10_000)
    args = parser.parse_args()

    path = Path(args.docx)
    doc = Document(path)
    blocks = []
    table_no = 0
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = " ".join(block.text.split())
            if text:
                blocks.append(("P", block.style.name, text))
        else:
            table_no += 1
            blocks.append(("TABLE", str(table_no), f"{len(block.rows)}x{len(block.columns)}"))
            for row_no, row in enumerate(block.rows, start=1):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                blocks.append(("ROW", f"{table_no}.{row_no}", " | ".join(cells)))

    print(
        f"FILE={path}\n"
        f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} "
        f"SECTIONS={len(doc.sections)} INLINE_SHAPES={len(doc.inline_shapes)} "
        f"BLOCK_LINES={len(blocks)}"
    )
    for idx, (kind, meta, text) in enumerate(blocks):
        if args.start <= idx < args.end:
            print(f"{idx:04d} [{kind}:{meta}] {text}")


if __name__ == "__main__":
    main()
