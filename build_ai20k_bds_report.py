from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = ROOT / "doc_assets"
OUT_PATH = OUT_DIR / "AI20K_BDS_O2O_Tai_lieu_tham_khao_va_de_xuat.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "64748B"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8F1F8"
PALE_GREEN = "EAF6EF"
PALE_GOLD = "FFF7DF"
PALE_RED = "FDECEC"
BORDER = "D8DEE8"
WHITE = "FFFFFF"
BLACK = "111827"

SOURCES = [
    (
        "ShowingTime+",
        "The Essential Guide to Real Estate Showings for New Agents",
        "https://showingtimeplus.com/hubfs/21555945/Gated%20Content/Guides/"
        "ShowingTimePlus_EssentialGuideToRealEstateShowingsForNewAgents.pdf",
    ),
    (
        "ShowingTime",
        "ShowingTime Appointment Center",
        "https://showingtimeplus.com/hubfs/uservoice/resources/general/"
        "ShowingTime_Appointment_Center.pdf",
    ),
    (
        "Google for Developers",
        "Google Calendar API - Freebusy: query",
        "https://developers.google.com/workspace/calendar/api/v3/reference/freebusy/query",
    ),
    (
        "Google for Developers",
        "Google Calendar API - Create events",
        "https://developers.google.com/workspace/calendar/api/guides/create-events",
    ),
    (
        "Microsoft Learn",
        "Microsoft Graph - calendar: getSchedule",
        "https://learn.microsoft.com/en-us/graph/api/calendar-getschedule?view=graph-rest-1.0",
    ),
    (
        "OpenAI Developers",
        "Function calling",
        "https://developers.openai.com/api/docs/guides/function-calling",
    ),
    (
        "LangChain",
        "LangGraph - Interrupts",
        "https://docs.langchain.com/oss/python/langgraph/interrupts",
    ),
    (
        "PostgreSQL",
        "PostgreSQL 18 Documentation - Explicit Locking",
        "https://www.postgresql.org/docs/current/explicit-locking.html",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            set_cell_width(cell, width)
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=11, color=BLACK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Calibri")
    r_pr.append(r_fonts)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        r_pr.append(underline_node)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), "20")
    r_pr.append(size_node)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def create_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId")))
        for el in numbering.findall(qn("w:abstractNum"))
        if el.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(el.get(qn("w:numId")))
        for el in numbering.findall(qn("w:num"))
        if el.get(qn("w:numId")) is not None
    ]

    def add_definition(kind: str, abstract_id: int, num_id: int) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)

    next_abs = max(abstract_ids or [0]) + 1
    next_num = max(num_ids or [0]) + 1
    add_definition("bullet", next_abs, next_num)
    bullet_id = next_num
    add_definition("decimal", next_abs + 1, next_num + 1)
    return bullet_id, next_num + 1


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_body(doc, text: str, refs: list[int] | None = None, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Normal")
    set_paragraph_spacing(p, after=6, line=1.10)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    if refs:
        ref_run = p.add_run(" " + "".join(f"[{n}]" for n in refs))
        set_run_font(ref_run, color=BLUE, bold=True)
    return p


def add_bullet(doc, text: str, bullet_id: int, refs: list[int] | None = None):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_id)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run)
    if refs:
        ref_run = p.add_run(" " + "".join(f"[{n}]" for n in refs))
        set_run_font(ref_run, color=BLUE, bold=True)
    return p


def add_number(doc, text: str, number_id: int, refs: list[int] | None = None):
    p = doc.add_paragraph()
    apply_numbering(p, number_id)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run)
    if refs:
        ref_run = p.add_run(" " + "".join(f"[{n}]" for n in refs))
        set_run_font(ref_run, color=BLUE, bold=True)
    return p


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_callout(doc, label: str, text: str, fill=PALE_BLUE, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.05)
    set_paragraph_spacing(p, before=5, after=9, line=1.10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, bold=True, color=INK)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=BLACK)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(label)
        set_run_font(r, size=9.5, color=INK, bold=True)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if len(value) <= 18 and idx != 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_paragraph_spacing(p, after=0, line=1.05)
            r = p.add_run(value)
            set_run_font(r, size=9.25, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(2)
    return table


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=3, after=8, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return p


def find_font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def pilc(value: str) -> str:
    return value if value.startswith("#") else f"#{value}"


def rounded_box(draw, xy, fill, outline, title, body, number=None):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=pilc(fill), outline=pilc(outline), width=3)
    if number is not None:
        draw.ellipse((x1 + 16, y1 + 15, x1 + 58, y1 + 57), fill=pilc(outline))
        num_font = find_font(21, bold=True)
        txt = str(number)
        bb = draw.textbbox((0, 0), txt, font=num_font)
        draw.text(
            (x1 + 37 - (bb[2] - bb[0]) / 2, y1 + 36 - (bb[3] - bb[1]) / 2 - 1),
            txt,
            font=num_font,
            fill=pilc("FFFFFF"),
        )
        title_x = x1 + 70
    else:
        title_x = x1 + 20
    title_font = find_font(25, bold=True)
    body_font = find_font(20, bold=False)
    draw.text((title_x, y1 + 16), title, font=title_font, fill=pilc("0B2545"))
    lines = body.split("\n")
    for idx, line in enumerate(lines):
        draw.text((x1 + 20, y1 + 67 + idx * 27), line, font=body_font, fill=pilc("334155"))


def arrow(draw, p1, p2, color="7C8EA3", width=4):
    draw.line([p1, p2], fill=pilc(color), width=width)
    x2, y2 = p2
    x1, y1 = p1
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - sign * 16, y2 - 9), (x2 - sign * 16, y2 + 9)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 9, y2 - sign * 16), (x2 + 9, y2 - sign * 16)]
    draw.polygon(points, fill=pilc(color))


def create_workflow_diagram(path: Path):
    img = Image.new("RGB", (1800, 930), pilc("FFFFFF"))
    d = ImageDraw.Draw(img)
    title_font = find_font(36, bold=True)
    d.text((70, 32), "Luồng đặt lịch có HITL và soft-hold", font=title_font, fill=pilc(INK))
    boxes = [
        (70, 115, 445, 280, "Tiếp nhận", "Khách mô tả nhu cầu\nbằng ngôn ngữ tự nhiên"),
        (515, 115, 890, 280, "Đối soát", "Lọc căn phù hợp và\nkiểm tra lịch sale"),
        (960, 115, 1335, 280, "Đề xuất", "Xếp hạng 2-3 khung giờ\nkèm căn có thể xem"),
        (1405, 115, 1780, 280, "Khách chọn", "Khách chọn giờ/căn\nvà xác nhận liên hệ"),
        (1405, 520, 1780, 685, "Sale duyệt", "Phê duyệt, chỉnh sửa\nhoặc từ chối đề xuất"),
        (960, 520, 1335, 685, "Khóa giao dịch", "Kiểm tra lại, tạo lịch\nvà soft-hold nguyên tử"),
        (515, 520, 890, 685, "Đồng bộ", "Tạo Calendar event,\ngửi xác nhận và nhắc"),
        (70, 520, 445, 685, "Theo dõi", "Dời/hủy, hết hạn hold,\nđo no-show và audit"),
    ]
    colors = ["E8F1F8", "EAF6EF", "FFF7DF", "F3EEFF", "FFF7DF", "FDECEC", "E8F1F8", "F2F4F7"]
    outlines = [BLUE, "2F855A", "A16B00", "6B46C1", "A16B00", "B42318", BLUE, "64748B"]
    for idx, (box, fill, outline) in enumerate(zip(boxes, colors, outlines), start=1):
        x1, y1, x2, y2, title, body = box
        rounded_box(d, (x1, y1, x2, y2), fill, outline, title, body, idx)
    arrow(d, (445, 198), (505, 198))
    arrow(d, (890, 198), (950, 198))
    arrow(d, (1335, 198), (1395, 198))
    arrow(d, (1592, 280), (1592, 510))
    arrow(d, (1405, 602), (1345, 602))
    arrow(d, (960, 602), (900, 602))
    arrow(d, (515, 602), (455, 602))
    note_font = find_font(23, bold=True)
    small_font = find_font(20)
    d.rounded_rectangle(
        (410, 770, 1390, 875),
        radius=16,
        fill=pilc("F8FAFC"),
        outline=pilc("CBD5E1"),
        width=2,
    )
    d.text((450, 790), "Điểm kiểm soát:", font=note_font, fill=pilc(INK))
    d.text(
        (650, 790),
        "Không giữ transaction trong lúc chờ sale; khi sale duyệt phải re-check và khóa lại.",
        font=small_font,
        fill=pilc("334155"),
    )
    d.text(
        (650, 830),
        "Chỉ thông báo “đã chốt” sau khi trạng thái nội bộ và đồng bộ lịch đạt điều kiện xác nhận.",
        font=small_font,
        fill=pilc("334155"),
    )
    img.save(path, quality=95)


def create_architecture_diagram(path: Path):
    img = Image.new("RGB", (1800, 1160), pilc("FFFFFF"))
    d = ImageDraw.Draw(img)
    title_font = find_font(36, bold=True)
    d.text((70, 30), "Kiến trúc đề xuất theo lớp", font=title_font, fill=pilc(INK))
    layers = [
        (
            100,
            120,
            1700,
            245,
            "Kênh người dùng",
            ["Next.js Web", "Khách hàng", "Sale / Điều phối", "WebSocket cập nhật"],
            "E8F1F8",
            BLUE,
        ),
        (
            100,
            300,
            1700,
            455,
            "API & điều phối nghiệp vụ",
            ["FastAPI", "Auth + RBAC", "Booking Service", "Approval Service", "Notification Service"],
            "EAF6EF",
            "2F855A",
        ),
        (
            100,
            510,
            1700,
            675,
            "AI Agent / Tool Orchestration",
            ["Responses + function tools", "LangGraph state", "HITL interrupt", "Policy & validation"],
            "FFF7DF",
            "A16B00",
        ),
        (
            100,
            730,
            1700,
            885,
            "Tích hợp ngoài",
            ["Google Calendar", "Microsoft Graph", "Email / SMS", "Bản đồ (nâng cao)"],
            "F3EEFF",
            "6B46C1",
        ),
        (
            100,
            940,
            1700,
            1125,
            "Dữ liệu & vận hành",
            ["PostgreSQL", "Soft-hold + audit", "Outbox / worker", "Logs + metrics"],
            "F2F4F7",
            "64748B",
        ),
    ]
    label_font = find_font(24, bold=True)
    item_font = find_font(21, bold=True)
    for x1, y1, x2, y2, label, items, fill, outline in layers:
        d.rounded_rectangle(
            (x1, y1, x2, y2),
            radius=20,
            fill=pilc(fill),
            outline=pilc(outline),
            width=3,
        )
        d.text((x1 + 25, y1 + 18), label, font=label_font, fill=pilc(INK))
        content_top = y1 + 62
        gap = (x2 - x1 - 90) / len(items)
        for idx, item in enumerate(items):
            bx1 = x1 + 25 + idx * gap
            bx2 = x1 + 25 + (idx + 1) * gap - 18
            d.rounded_rectangle(
                (bx1, content_top, bx2, y2 - 18),
                radius=12,
                fill=pilc("FFFFFF"),
                outline=pilc(outline),
                width=2,
            )
            bb = d.textbbox((0, 0), item, font=item_font)
            tx = (bx1 + bx2) / 2 - (bb[2] - bb[0]) / 2
            ty = content_top + (y2 - 18 - content_top) / 2 - (bb[3] - bb[1]) / 2
            d.text((tx, ty), item, font=item_font, fill=pilc("334155"))
    for y_start, y_end in [(245, 300), (455, 510), (675, 730), (885, 940)]:
        arrow(d, (900, y_start + 8), (900, y_end - 8), color="94A3B8", width=5)
    img.save(path, quality=95)


def add_picture_with_alt(doc, path: Path, width_inches: float, alt_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=2, line=1.0)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    return shape


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color, before, after in [
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(hp, after=0, line=1.0)
    r1 = hp.add_run("AI20K  |  BĐS O2O - AI Agent đặt lịch xem nhà")
    set_run_font(r1, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(fp, after=0, line=1.0)
    r2 = fp.add_run("Tài liệu tham khảo & đề xuất  •  Trang ")
    set_run_font(r2, size=8.5, color=MUTED)
    add_page_field(fp)


def add_cover(doc: Document):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(64)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(kicker, after=16, line=1.0)
    r = kicker.add_run("AI20K  •  NGÂN HÀNG ĐỀ KHÓA 3 & 4")
    set_run_font(r, size=11, color="A16B00", bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, after=10, line=1.0)
    r = title.add_run("BĐS - KINH DOANH O2O")
    set_run_font(r, size=28, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=8, line=1.05)
    r = subtitle.add_run("AI Agent điều phối lịch xem nhà và giữ căn tạm thời")
    set_run_font(r, size=16, color=DARK_BLUE, bold=True)

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.left_indent = Inches(0.65)
    desc.paragraph_format.right_indent = Inches(0.65)
    set_paragraph_spacing(desc, after=38, line=1.15)
    r = desc.add_run(
        "Tài liệu tham khảo giải pháp bên ngoài, phân tích bài toán và đề xuất "
        "kiến trúc triển khai cho Doanh nghiệp bất động sản X"
    )
    set_run_font(r, size=11.5, color=MUTED, italic=True)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(line, after=32, line=1.0)
    r = line.add_run("—  TÀI LIỆU NGHIÊN CỨU & KỊCH BẢN DEMO  —")
    set_run_font(r, size=10, color="A16B00", bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(meta, after=4, line=1.0)
    r = meta.add_run("Phiên bản 1.0  |  28/07/2026")
    set_run_font(r, size=10.5, color=INK, bold=True)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note, after=0, line=1.0)
    r = note.add_run("Phục vụ tổng hợp bài tập lớn và trình diễn MVP")
    set_run_font(r, size=9.5, color=MUTED)

    doc.add_page_break()


def build_document():
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)
    workflow_path = ASSET_DIR / "workflow_hitl_softhold.png"
    architecture_path = ASSET_DIR / "architecture_layers.png"
    create_workflow_diagram(workflow_path)
    create_architecture_diagram(architecture_path)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bullet_id, number_id = create_numbering(doc)
    add_cover(doc)

    add_heading(doc, "Tóm tắt điều hành", 1)
    add_body(
        doc,
        "Đề tài giải quyết một nút thắt rất thực tế: khách muốn xem nhà nhưng lịch sale, "
        "xe đưa đón, trạng thái căn và khả năng phục vụ tại điểm xem đang được điều phối rời rạc "
        "qua chat. Hệ quả là phản hồi chậm, trùng lịch, giữ căn không kiểm soát và khó truy vết.",
    )
    add_callout(
        doc,
        "Khuyến nghị trọng tâm",
        "Xây MVP theo mô hình “AI đề xuất - con người phê duyệt - hệ thống khóa giao dịch”. "
        "LLM chỉ hiểu ý định và chọn công cụ; quyền quyết định chốt lịch/giữ căn nằm ở sale và "
        "các quy tắc giao dịch trong backend.",
        fill=PALE_BLUE,
        border=BLUE,
    )
    add_body(
        doc,
        "Các hệ thống quản lý lịch xem bất động sản đã chứng minh giá trị của đặt lịch đa kênh, "
        "phê duyệt lịch, đồng bộ lịch cá nhân và theo dõi toàn bộ hoạt động. Từ đó, nhóm có thể "
        "tập trung phần đổi mới vào hội thoại tự nhiên, phối hợp nhiều nguồn lực, soft-hold có "
        "hạn và xử lý ngoại lệ tự động.",
        refs=[1, 2],
    )
    add_heading(doc, "Kết quả nên đạt trong bản demo", 2)
    for item in [
        "Khách nhập yêu cầu tự nhiên, ví dụ: “Chiều thứ Bảy tôi muốn xem căn 2 phòng ngủ ở khu Đông, có xe đón từ quận 1.”",
        "Agent gọi công cụ để lọc căn, kiểm tra lịch sale và đề xuất 2-3 lựa chọn có giải thích.",
        "Khách chọn một phương án; sale nhận thẻ phê duyệt và có thể duyệt, sửa hoặc từ chối.",
        "Sau khi duyệt, backend tái kiểm tra và tạo soft-hold trong giao dịch ngắn; lịch được đồng bộ và hai bên nhận xác nhận.",
        "Demo được ít nhất một ngoại lệ: xung đột lịch, căn vừa bị giữ, API Calendar lỗi hoặc hold hết hạn.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "Cấu trúc tài liệu", 2)
    structure_rows = [
        ["1-2", "Bối cảnh, mục tiêu và kinh nghiệm từ giải pháp bên ngoài"],
        ["3-5", "Quy trình nghiệp vụ, phạm vi MVP và kiến trúc đề xuất"],
        ["6-8", "Thiết kế Agent, HITL, soft-hold và chống double-booking"],
        ["9-11", "Lỗi, bảo mật, kịch bản demo, đánh giá và lộ trình"],
        ["12", "Kết luận, checklist bàn giao và nguồn tham khảo"],
    ]
    add_table(doc, ["Phần", "Nội dung"], structure_rows, [1500, 7860])

    add_heading(doc, "1. Phân tích bài toán", 1)
    add_heading(doc, "1.1. Thực trạng và nguyên nhân gốc", 2)
    add_body(
        doc,
        "Quy trình hiện tại phụ thuộc vào trao đổi nối tiếp giữa khách, sale, điều phối viên "
        "và bộ phận quản lý căn. Mỗi bên giữ một phần dữ liệu; lịch trong chat có thể không còn "
        "đúng tại thời điểm chốt. Đây không chỉ là bài toán chatbot mà là bài toán điều phối "
        "tài nguyên có trạng thái và hành động tạo tác động thật.",
    )
    cause_rows = [
        ["Thông tin phân tán", "Chat, lịch cá nhân, file căn, điện thoại", "Khó xác định dữ liệu nào mới nhất"],
        ["Kiểm tra nối tiếp", "Hỏi sale rồi mới hỏi căn/xe", "Thời gian phản hồi dài, khách rời bỏ"],
        ["Không có khóa", "Hai luồng cùng thấy một căn còn trống", "Double-booking hoặc giữ căn chồng chéo"],
        ["Thiếu trạng thái chuẩn", "“Đang hỏi”, “ok”, “giữ nhé” trong chat", "Khó audit, khó tự động dời/hủy"],
        ["Phụ thuộc API ngoài", "Calendar/SMS/Inventory có thể timeout", "Dễ xác nhận sai hoặc gửi trùng"],
    ]
    add_table(doc, ["Nguyên nhân", "Biểu hiện", "Hệ quả"], cause_rows, [1900, 3400, 4060])

    add_heading(doc, "1.2. Mục tiêu nghiệp vụ", 2)
    for item in [
        "Rút ngắn thời gian từ lúc khách nêu nhu cầu đến lúc nhận được phương án khả thi.",
        "Tăng tỷ lệ đặt lịch thành công nhưng không đánh đổi độ chính xác hoặc quyền kiểm soát của sale.",
        "Ngăn trùng lịch sale và trùng giữ căn bằng ràng buộc dữ liệu, không dựa vào trí nhớ của Agent.",
        "Cho phép dời/hủy, nhắc lịch và giải phóng hold theo cùng một quy trình có audit.",
        "Thu thập dữ liệu để đo tỷ lệ xác nhận, no-show, thời gian phê duyệt và nguyên nhân thất bại.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "1.3. Nguyên tắc thiết kế", 2)
    principles = [
        ["LLM không phải nguồn sự thật", "Lịch và trạng thái căn luôn lấy từ tool/backend tại thời điểm thao tác."],
        ["Đọc tự động, ghi có kiểm soát", "Tra cứu có thể tự động; chốt lịch và giữ căn phải qua HITL."],
        ["Tái kiểm tra trước khi ghi", "Dữ liệu có thể thay đổi giữa lúc đề xuất và lúc sale duyệt."],
        ["Giao dịch ngắn", "Không giữ row-lock trong lúc chờ người dùng/sale phản hồi."],
        ["Idempotent", "Retry không được tạo lịch, hold hoặc thông báo trùng."],
    ]
    add_table(doc, ["Nguyên tắc", "Áp dụng"], principles, [2500, 6860])

    doc.add_page_break()
    add_heading(doc, "2. Tham khảo cách các hệ thống bên ngoài giải quyết", 1)
    add_body(
        doc,
        "Phần này không nhằm sao chép một sản phẩm cụ thể. Mục tiêu là rút ra các mẫu thiết kế "
        "đã được dùng trong quản lý lịch xem, lịch doanh nghiệp và hệ thống Agent, sau đó điều "
        "chỉnh cho yêu cầu HITL và soft-hold của đề tài.",
    )
    benchmark_rows = [
        [
            "ShowingTime+",
            "Yêu cầu lịch xem; bên quản lý/listing agent xin phê duyệt; trao đổi qua điện thoại, tin nhắn, email hoặc ứng dụng.",
            "Tách yêu cầu khỏi xác nhận; hỗ trợ đa kênh và theo dõi lịch sử.",
        ],
        [
            "ShowingTime Appointment Center",
            "Đặt lịch nhanh, đồng bộ lịch cá nhân, thông báo nhiều kênh, theo dõi showing/feedback.",
            "MVP nên có timeline sự kiện và trạng thái đồng bộ rõ ràng.",
        ],
        [
            "Google Calendar",
            "FreeBusy để xem bận/rảnh; Events để tạo lịch, khách mời, địa điểm và nhắc lịch.",
            "Chỉ xin quyền cần thiết; dùng event ID/idempotency để tránh tạo trùng.",
        ],
        [
            "Microsoft Graph",
            "getSchedule lấy availability của nhiều người hoặc tài nguyên theo một khoảng thời gian.",
            "Thiết kế adapter lịch chung để đổi Google/Outlook mà không đổi logic Agent.",
        ],
        [
            "LangGraph",
            "Interrupt lưu trạng thái, tạm dừng và tiếp tục sau quyết định con người.",
            "Phù hợp thẻ Approve/Edit/Reject; tác vụ trước interrupt phải idempotent.",
        ],
        [
            "PostgreSQL",
            "FOR UPDATE chặn các giao dịch khác khóa hoặc sửa cùng hàng cho đến khi transaction kết thúc.",
            "Dùng khóa trong giao dịch ngắn; lưu hold thành bản ghi có expires_at.",
        ],
    ]
    add_table(doc, ["Nguồn", "Cách họ làm", "Bài học áp dụng"], benchmark_rows, [1850, 4050, 3460])
    add_body(
        doc,
        "ShowingTime+ mô tả lịch xem là một cuộc hẹn chuyên nghiệp giữa các bên, trong đó yêu "
        "cầu ngày/giờ cần được bên liên quan chấp thuận. Tài liệu Appointment Center nhấn mạnh "
        "đồng bộ lịch cá nhân, thông báo đa kênh và theo dõi hoạt động.",
        refs=[1, 2],
    )
    add_body(
        doc,
        "Google Calendar cung cấp FreeBusy cho một tập lịch và cho phép tạo sự kiện với thời "
        "gian, khách mời, địa điểm và nhắc lịch. Tài liệu Google cũng khuyến nghị ứng dụng có "
        "thể tự đặt event ID phù hợp để hỗ trợ đồng bộ và tránh tạo trùng sau lỗi mạng.",
        refs=[3, 4],
    )
    add_body(
        doc,
        "Microsoft Graph getSchedule trả về thông tin rảnh/bận cho nhiều người hoặc tài nguyên "
        "trong một khoảng xác định; vì vậy cùng một miền nghiệp vụ có thể được che bởi interface "
        "CalendarProvider thay vì viết logic riêng trong Agent.",
        refs=[5],
    )
    add_body(
        doc,
        "OpenAI function calling là vòng lặp nhiều bước: ứng dụng khai báo tools, mô hình đề "
        "xuất tool call, ứng dụng tự thực thi, gửi kết quả trở lại rồi mô hình mới tổng hợp phản "
        "hồi. Lược đồ strict giúp đối số tool có cấu trúc chặt hơn, nhưng backend vẫn phải xác "
        "thực quyền và quy tắc nghiệp vụ.",
        refs=[6],
    )
    add_body(
        doc,
        "LangGraph interrupt cho phép lưu checkpoint, dừng tại điểm cần con người và tiếp tục "
        "bằng cùng thread. Do node có thể chạy lại từ đầu khi resume, mọi side effect trước điểm "
        "interrupt phải có idempotency key hoặc được chuyển sang bước sau phê duyệt.",
        refs=[7],
    )

    add_heading(doc, "2.1. Khoảng trống mà đề tài có thể tạo khác biệt", 2)
    for item in [
        "Hiểu câu tự nhiên có nhiều ràng buộc cùng lúc: khu vực, loại căn, thời gian, điểm đón và người đi cùng.",
        "Điều phối đồng thời sale, căn, xe và phòng chờ thay vì chỉ một lịch cá nhân.",
        "Giải thích vì sao đề xuất một khung giờ và tự đưa phương án thay thế khi có xung đột.",
        "Kết hợp HITL với khóa giao dịch và soft-hold có hạn, thay vì chỉ gửi một yêu cầu chờ xác nhận.",
        "Đo chất lượng Agent bằng kết quả nghiệp vụ, không chỉ độ “hay” của câu trả lời.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_page_break()
    add_heading(doc, "3. Quy trình nghiệp vụ đề xuất", 1)
    add_picture_with_alt(
        doc,
        workflow_path,
        6.35,
        "Sơ đồ 8 bước đặt lịch xem nhà: tiếp nhận, đối soát, đề xuất, khách chọn, sale duyệt, khóa giao dịch, đồng bộ và theo dõi.",
    )
    add_caption(doc, "Hình 1. Luồng đặt lịch có HITL và soft-hold")
    add_heading(doc, "3.1. Luồng chuẩn", 2)
    steps = [
        "Agent trích xuất nhu cầu và hỏi lại phần còn thiếu: ngày/khung giờ, khu vực, loại căn, số người, điểm đón và kênh liên hệ.",
        "Backend lọc danh sách căn phù hợp, chỉ trả cho Agent dữ liệu tối thiểu cần thiết.",
        "Calendar adapter lấy FreeBusy của các sale phù hợp; resource service kiểm tra xe và phòng chờ nếu có.",
        "Scheduler tính giao của các khoảng trống, chèn buffer di chuyển và xếp hạng 2-3 phương án.",
        "Khách chọn phương án; hệ thống tạo Appointment ở trạng thái PENDING_SALE_APPROVAL, chưa giữ căn.",
        "Sale nhận thẻ duyệt với đầy đủ thông tin và chọn Approve, Edit hoặc Reject.",
        "Khi Approve, backend tái kiểm tra và trong một transaction ngắn: khóa dữ liệu liên quan, ghi appointment, tạo soft-hold và outbox event.",
        "Worker đồng bộ Calendar/thông báo. Khi đạt điều kiện xác nhận, giao diện cập nhật real-time cho khách và sale.",
    ]
    for step in steps:
        add_number(doc, step, number_id)

    add_heading(doc, "3.2. Vì sao không giữ căn trước khi sale duyệt?", 2)
    add_body(
        doc,
        "Ràng buộc đề bài yêu cầu sale xác nhận trước khi chốt lịch và giữ căn. Nếu Agent tạo "
        "hold ngay khi khách mới chọn, nhiều căn có thể bị chiếm bởi yêu cầu chưa được kiểm tra. "
        "Nếu giữ row-lock trong lúc đợi sale, transaction kéo dài gây nghẽn và tăng nguy cơ deadlock.",
        refs=[8],
    )
    add_callout(
        doc,
        "Thiết kế đúng",
        "Trước phê duyệt chỉ lưu “ý định đặt lịch”. Sau phê duyệt mới re-check và tạo soft-hold "
        "bền vững trong database. Nếu căn đã bị giữ trong thời gian chờ, hệ thống trả 409 CONFLICT "
        "và đề xuất phương án gần nhất.",
        fill=PALE_GREEN,
        border="2F855A",
    )

    add_heading(doc, "3.3. Máy trạng thái cuộc hẹn", 2)
    states = [
        ["DRAFT", "Đang hội thoại, chưa đủ dữ liệu", "Khách/Agent"],
        ["OPTIONS_PROPOSED", "Đã có các phương án khả thi", "Agent"],
        ["PENDING_SALE_APPROVAL", "Khách đã chọn, chờ sale", "Khách"],
        ["RESERVING", "Đang khóa và ghi dữ liệu", "Backend"],
        ["CONFIRMED", "Đã giữ căn và đồng bộ đủ điều kiện", "Backend"],
        ["RESCHEDULE_PENDING", "Khách/sale yêu cầu đổi", "Khách/Sale"],
        ["CANCELLED", "Đã hủy và giải phóng tài nguyên", "Backend"],
        ["EXPIRED", "Quá hạn phê duyệt/soft-hold", "Scheduler"],
        ["SYNC_FAILED", "Dữ liệu nội bộ có nhưng tích hợp ngoài lỗi", "Worker"],
    ]
    add_table(doc, ["Trạng thái", "Ý nghĩa", "Chủ thể kích hoạt"], states, [2600, 4800, 1960])

    doc.add_page_break()
    add_heading(doc, "4. Phạm vi chức năng", 1)
    add_heading(doc, "4.1. Vai trò", 2)
    roles = [
        ["Khách hàng", "Nêu nhu cầu, chọn phương án, xác nhận thông tin, yêu cầu dời/hủy."],
        ["Sale", "Quản lý lịch cá nhân, duyệt/sửa/từ chối, xem chi tiết khách được phân công."],
        ["Điều phối viên", "Xem toàn bộ luồng, can thiệp lỗi, phân sale/xe, xử lý tình huống đặc biệt."],
        ["Quản trị viên", "Quản lý người dùng, quyền, cấu hình hold, kênh thông báo và audit."],
        ["AI Agent", "Hiểu ý định, gọi tool, giải thích và dẫn hội thoại; không tự bỏ qua policy."],
    ]
    add_table(doc, ["Vai trò", "Quyền/trách nhiệm chính"], roles, [1900, 7460])

    add_heading(doc, "4.2. Chức năng cơ bản bắt buộc", 2)
    for item in [
        "Đăng nhập và phân quyền tối thiểu: khách, sale/điều phối và quản trị.",
        "Hội thoại đặt lịch có lưu thread và trạng thái nghiệp vụ.",
        "Tìm căn, kiểm tra trạng thái căn và lịch sale.",
        "Đề xuất nhiều khung giờ có giải thích ngắn gọn.",
        "Thẻ HITL cho sale: Approve, Edit, Reject; mọi quyết định có audit.",
        "Tạo soft-hold có expires_at, cảnh báo sắp hết hạn và job giải phóng.",
        "Đồng bộ Calendar, gửi xác nhận/nhắc, hỗ trợ dời và hủy.",
        "WebSocket hoặc Server-Sent Events để cập nhật trạng thái ngay trên giao diện.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "4.3. Chức năng nâng cao", 2)
    advanced = [
        ["Tối ưu lộ trình", "Xếp lịch nhiều khách theo khoảng cách, thời gian di chuyển và buffer."],
        ["Tự động dời lịch", "Phát hiện xung đột mới, đề xuất lại và xin phê duyệt thay đổi."],
        ["Bản đồ tiện ích", "Hiển thị trường học, bệnh viện, giao thông và thời gian đến điểm xem."],
        ["Memory sở thích", "Lưu khung giờ, khu vực, loại căn ưa thích với sự đồng ý của khách."],
        ["Eval liên tục", "Theo dõi booking success, no-show, tool error, conflict và latency."],
    ]
    add_table(doc, ["Hạng mục", "Giá trị"], advanced, [2200, 7160])

    add_heading(doc, "4.4. Ngoài phạm vi MVP", 2)
    add_body(
        doc,
        "Không nên đưa vào demo đầu: thanh toán/đặt cọc thật, ký hợp đồng, định giá tự động, "
        "quyết định ưu tiên khách theo thuộc tính nhạy cảm, hoặc tối ưu tuyến xe quy mô lớn. "
        "Các phần này làm tăng rủi ro và che mất luồng cốt lõi cần chứng minh.",
    )

    add_heading(doc, "5. Kiến trúc kỹ thuật đề xuất", 1)
    add_picture_with_alt(
        doc,
        architecture_path,
        6.35,
        "Sơ đồ kiến trúc năm lớp gồm giao diện, API nghiệp vụ, AI Agent, tích hợp ngoài và dữ liệu vận hành.",
    )
    add_caption(doc, "Hình 2. Kiến trúc đề xuất theo lớp")
    add_heading(doc, "5.1. Trách nhiệm từng thành phần", 2)
    components = [
        ["Next.js", "Chat, danh sách phương án, thẻ HITL, dashboard điều phối và cập nhật real-time."],
        ["FastAPI", "Auth/RBAC, API nghiệp vụ, validation, transaction, idempotency và WebSocket."],
        ["Agent Orchestrator", "Giữ state hội thoại, chọn tool, HITL interrupt, retry có giới hạn."],
        ["Calendar Adapter", "Chuẩn hóa Google FreeBusy/Events và Microsoft getSchedule."],
        ["Booking Service", "Máy trạng thái, kiểm tra xung đột, soft-hold, dời/hủy."],
        ["PostgreSQL", "Nguồn sự thật cho căn, lịch nội bộ, hold, outbox, audit và memory."],
        ["Worker/Outbox", "Đồng bộ API ngoài, retry, gửi nhắc và giải phóng hold hết hạn."],
    ]
    add_table(doc, ["Thành phần", "Trách nhiệm"], components, [2150, 7210])

    add_heading(doc, "5.2. Quyết định kiến trúc quan trọng", 2)
    for item in [
        "Agent không truy cập database trực tiếp; mọi thao tác đi qua tool có schema, auth context và validation.",
        "Lịch nội bộ là lớp bảo vệ chống xung đột; Calendar ngoài là kênh đồng bộ, không phải khóa duy nhất.",
        "Dùng outbox để tách transaction nội bộ khỏi API ngoài; worker retry không làm trùng nhờ idempotency key.",
        "WebSocket chỉ phát sự kiện trạng thái; client vẫn tải lại từ API nếu mất kết nối.",
        "Triển khai Fly.io có thể tách web/API/worker, nhưng MVP có thể chạy Docker Compose để demo ổn định.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_page_break()
    add_heading(doc, "6. Thiết kế AI Agent và bộ công cụ", 1)
    add_heading(doc, "6.1. Ranh giới trách nhiệm", 2)
    boundary_rows = [
        ["LLM/Agent", "Hiểu câu, hỏi bổ sung, chọn tool, so sánh phương án, diễn đạt kết quả."],
        ["Backend", "Kiểm tra quyền, dữ liệu thật, xung đột, transaction, hold, idempotency, audit."],
        ["Con người", "Duyệt/sửa/từ chối hành động chốt lịch và giữ căn."],
    ]
    add_table(doc, ["Lớp", "Được phép quyết định"], boundary_rows, [1900, 7460])
    add_body(
        doc,
        "Function calling cho phép mô hình đề xuất gọi hàm có đối số JSON, nhưng code ứng dụng "
        "mới là nơi thực thi và trả kết quả. Do đó có thể thay mô hình mà không thay các quy tắc "
        "giao dịch cốt lõi.",
        refs=[6],
    )

    add_heading(doc, "6.2. Danh mục tools tối thiểu", 2)
    tools_rows = [
        ["search_units", "Đọc", "Lọc căn theo tiêu chí; không trả dữ liệu nhạy cảm."],
        ["check_unit_availability", "Đọc", "Kiểm tra trạng thái hiện tại và khoảng xem."],
        ["check_sale_availability", "Đọc", "Lấy FreeBusy qua Calendar adapter."],
        ["rank_visit_options", "Đọc", "Tạo/xếp hạng phương án có buffer."],
        ["request_sale_approval", "Ghi nhẹ", "Tạo pending request; chưa giữ căn."],
        ["create_soft_hold", "Ghi nguy hiểm", "Chỉ chạy sau approval token hợp lệ."],
        ["sync_calendar_event", "Ghi ngoài", "Idempotent; chạy qua outbox/worker."],
        ["reschedule_appointment", "Ghi nguy hiểm", "Tái kiểm tra, cần HITL theo policy."],
        ["cancel_appointment", "Ghi nguy hiểm", "Hủy lịch, release hold, gửi thông báo."],
    ]
    add_table(doc, ["Tool", "Loại", "Quy tắc"], tools_rows, [2650, 1350, 5360])

    add_heading(doc, "6.3. Quy ước schema và lỗi", 2)
    for item in [
        "Mỗi tool có JSON Schema chặt, required fields rõ và additionalProperties=false nếu phù hợp.",
        "Mọi tool nhận auth_context từ server, không tin role/user_id do mô hình tự truyền.",
        "Kết quả thống nhất: ok, data, error_code, retryable, user_message, trace_id.",
        "Tool ghi bắt buộc có idempotency_key; tool nguy hiểm cần approval_id/approval_version.",
        "Agent không tự retry vô hạn; tối đa 2-3 lần cho lỗi tạm thời rồi chuyển điều phối viên.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(
        doc,
        "Ví dụ policy",
        "check_sale_availability có thể tự chạy; create_soft_hold bị chặn nếu không có approval "
        "còn hiệu lực. Sale sửa giờ đồng nghĩa approval_version tăng và phải re-check toàn bộ.",
        fill=PALE_GOLD,
        border="A16B00",
    )

    add_heading(doc, "7. Dữ liệu, khóa giao dịch và soft-hold", 1)
    add_heading(doc, "7.1. Mô hình dữ liệu cốt lõi", 2)
    data_rows = [
        ["users", "id, role, status", "Tài khoản và phân quyền"],
        ["customers", "id, contact_ref, consent", "Thông tin khách đã tối thiểu hóa"],
        ["sales", "user_id, calendar_provider, service_area", "Sale và phạm vi phục vụ"],
        ["inventory_units", "id, project_id, status, version", "Nguồn sự thật trạng thái căn"],
        ["appointments", "id, customer_id, sale_id, unit_id, time_range, status", "Cuộc hẹn"],
        ["approvals", "id, appointment_id, version, decision, decided_by", "Bằng chứng HITL"],
        ["soft_holds", "id, unit_id, appointment_id, expires_at, status", "Giữ căn có hạn"],
        ["outbox_events", "id, aggregate_id, type, payload, status", "Đồng bộ tin cậy"],
        ["audit_logs", "actor, action, object, before, after, trace_id", "Truy vết"],
    ]
    add_table(doc, ["Bảng", "Trường chính", "Mục đích"], data_rows, [1900, 4300, 3160])

    add_heading(doc, "7.2. Thuật toán chống double-booking", 2)
    lock_steps = [
        "Nhận approval_id và kiểm tra approval còn hiệu lực, đúng phiên bản và đúng sale.",
        "BEGIN transaction; SELECT inventory_unit ... FOR UPDATE; khóa các tài nguyên theo thứ tự cố định.",
        "Kiểm tra lại căn, time_range của sale, xe/phòng chờ và soft-hold còn hiệu lực.",
        "Nếu xung đột: ROLLBACK, trả 409 CONFLICT cùng dữ liệu để Agent đề xuất lại.",
        "Nếu hợp lệ: ghi appointment, soft_hold, outbox_event và idempotency record.",
        "COMMIT ngay; không gọi Calendar/SMS bên trong transaction.",
        "Worker đọc outbox và gọi API ngoài; kết quả cập nhật trạng thái bằng một transaction riêng.",
    ]
    for step in lock_steps:
        add_number(doc, step, number_id)
    add_body(
        doc,
        "PostgreSQL xác định FOR UPDATE là khóa hàng ngăn giao dịch khác khóa, sửa hoặc xóa cùng "
        "hàng cho đến khi transaction kết thúc. Tài liệu cũng lưu ý deadlock và khuyến nghị "
        "các luồng lấy nhiều khóa theo cùng thứ tự.",
        refs=[8],
    )

    add_heading(doc, "7.3. Quy tắc soft-hold", 2)
    hold_rows = [
        ["Thời hạn mặc định", "15 phút sau sale duyệt (cấu hình theo dự án)."],
        ["Cảnh báo", "Gửi cảnh báo khi còn 5 phút; hiển thị đồng hồ đếm ngược."],
        ["Gia hạn", "Chỉ sale/điều phối được gia hạn; ghi lý do và giới hạn số lần."],
        ["Hết hạn", "Job chuyển EXPIRED, giải phóng hold và phát sự kiện real-time."],
        ["Chuyển trạng thái", "Soft-hold không đồng nghĩa đặt cọc/đã bán; phải hiển thị rõ."],
        ["Xung đột", "Không tự chiếm căn khác; Agent phải xin khách/sale chọn phương án."],
    ]
    add_table(doc, ["Quy tắc", "Đề xuất"], hold_rows, [2500, 6860])

    doc.add_page_break()
    add_heading(doc, "8. HITL, đồng bộ lịch và nhất quán phân tán", 1)
    add_heading(doc, "8.1. Trải nghiệm phê duyệt", 2)
    add_body(
        doc,
        "Thẻ phê duyệt cần cho sale thấy đủ dữ liệu để quyết định trong một màn hình: khách, "
        "căn, thời gian, địa điểm/điểm đón, buffer, thời hạn phản hồi và cảnh báo xung đột. Ba "
        "quyết định là Approve, Edit và Reject; mọi quyết định phải có timestamp, actor và version.",
    )
    add_body(
        doc,
        "LangGraph interrupt phù hợp vì có thể checkpoint trạng thái, tạm dừng và tiếp tục đúng "
        "thread sau khi sale phản hồi. Tuy nhiên, node có thể chạy lại khi resume nên các side "
        "effect trước interrupt phải idempotent hoặc được đưa sang node sau phê duyệt.",
        refs=[7],
    )
    add_heading(doc, "8.2. Saga đồng bộ Calendar", 2)
    saga_rows = [
        ["1. Commit nội bộ", "Appointment + hold + outbox", "Nguồn sự thật đã có phiên bản"],
        ["2. Tạo event", "Worker gọi Calendar với idempotency key", "Nhận external_event_id"],
        ["3. Gửi thông báo", "Email/SMS/push theo cấu hình", "Ghi notification status"],
        ["4. Xác nhận", "Chuyển CONFIRMED khi đủ điều kiện", "Phát WebSocket event"],
        ["5. Bù trừ", "Lỗi kéo dài: release/flag/manual review", "Không im lặng xác nhận sai"],
    ]
    add_table(doc, ["Bước", "Hành động", "Kết quả"], saga_rows, [1450, 4440, 3470])
    add_body(
        doc,
        "Google Calendar FreeBusy trả thông tin bận/rảnh cho nhiều lịch; Events API tạo lịch "
        "và có thể gửi cập nhật cho khách mời. Microsoft Graph cung cấp getSchedule cho người "
        "hoặc tài nguyên. Adapter nên trả cùng một cấu trúc nội bộ để Scheduler không phụ thuộc "
        "nhà cung cấp.",
        refs=[3, 4, 5],
    )
    add_heading(doc, "8.3. Điều kiện hiển thị “đã xác nhận”", 2)
    for item in [
        "Appointment nội bộ ở CONFIRMED hoặc CONFIRMED_WITH_SYNC_WARNING theo policy đã công bố.",
        "Soft-hold ACTIVE và expires_at còn hiệu lực.",
        "Không có xung đột sale/unit/resource tại version hiện tại.",
        "Nếu Calendar chưa đồng bộ, UI phải nói rõ “đã giữ nội bộ, đang đồng bộ lịch” thay vì báo hoàn tất.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "9. Xử lý lỗi và bảo mật", 1)
    add_heading(doc, "9.1. Ma trận lỗi", 2)
    error_rows = [
        ["Calendar timeout", "Retry exponential + jitter; cùng idempotency key", "Đang đồng bộ, chưa gửi lại yêu cầu"],
        ["Inventory không phản hồi", "Không tạo hold; circuit breaker; cảnh báo điều phối", "Chưa thể xác nhận trạng thái căn"],
        ["Căn bị giữ sau khi đề xuất", "409, tải lựa chọn gần nhất", "Căn vừa thay đổi; chọn phương án khác"],
        ["Sale không phản hồi", "Nhắc, hết hạn pending, chuyển sale nếu policy cho phép", "Yêu cầu đang chờ xác nhận"],
        ["Hold sắp hết hạn", "Cảnh báo 5 phút; yêu cầu gia hạn có audit", "Còn X phút giữ tạm"],
        ["Thông báo lỗi", "Retry riêng; không rollback booking đã hợp lệ", "Lịch vẫn giữ; kênh gửi đang thử lại"],
        ["WebSocket mất kết nối", "Client reconnect + GET trạng thái mới nhất", "Đang kết nối lại"],
    ]
    add_table(doc, ["Tình huống", "Xử lý hệ thống", "Thông điệp người dùng"], error_rows, [2100, 4320, 2940])

    add_heading(doc, "9.2. Kiểm soát bảo mật", 2)
    for item in [
        "RBAC và kiểm tra quyền ở backend cho từng object; sale chỉ xem khách/lịch được phân công.",
        "Tối thiểu hóa dữ liệu gửi cho LLM; thay số điện thoại/email bằng customer_ref khi không cần.",
        "Mã hóa TLS khi truyền; mã hóa/secret manager cho token Calendar và thông tin liên hệ.",
        "Không ghi prompt/tool log chứa dữ liệu nhạy cảm ở mức debug trên production; có cơ chế che dữ liệu.",
        "Audit các hành động đọc nhạy cảm, duyệt, giữ, gia hạn, dời và hủy.",
        "OAuth theo nguyên tắc quyền tối thiểu; tách token theo người dùng/nhà cung cấp.",
        "Rate limit, chống prompt injection qua dữ liệu listing và allowlist tools theo vai trò.",
        "Chính sách retention/consent cho memory sở thích; cho phép khách yêu cầu xóa.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(
        doc,
        "Nguyên tắc",
        "Agent có thể đề xuất hành động nhưng không được tự nâng quyền, tự tạo approval token "
        "hoặc coi dữ liệu trong prompt là bằng chứng đã phê duyệt.",
        fill=PALE_RED,
        border="B42318",
    )

    doc.add_page_break()
    add_heading(doc, "10. Kịch bản demo tối nay", 1)
    add_heading(doc, "10.1. Kịch bản chính (6-8 phút)", 2)
    demo_steps = [
        "Đăng nhập bằng tài khoản khách và nhập: “Chiều thứ Bảy tôi muốn xem căn 2 phòng ngủ ở khu Đông; đi 3 người, cần xe đón ở quận 1.”",
        "Agent hiển thị phần đã hiểu và gọi search_units, check_sale_availability, check resources.",
        "Màn hình trả ba phương án; mỗi phương án có căn, sale, giờ, điểm đón và lý do xếp hạng.",
        "Khách chọn phương án 14:00; trạng thái chuyển sang Chờ sale xác nhận.",
        "Chuyển sang tài khoản sale: mở thẻ HITL, chỉnh buffer từ 15 thành 20 phút rồi Approve.",
        "Màn hình khách nhận cập nhật real-time: đang khóa → đang đồng bộ → đã xác nhận; hiển thị thời hạn soft-hold.",
        "Mở lịch/khung mock Calendar để chứng minh sự kiện đã được tạo và thông báo đã được ghi.",
    ]
    for step in demo_steps:
        add_number(doc, step, number_id)

    add_heading(doc, "10.2. Ngoại lệ bắt buộc nên demo", 2)
    add_callout(
        doc,
        "Xung đột có chủ đích",
        "Trước khi sale duyệt, dùng tài khoản điều phối giữ cùng căn ở khung 14:00. Khi sale "
        "Approve, transaction trả 409; Agent không xác nhận sai mà đề xuất căn tương đương lúc "
        "14:30 hoặc cùng căn lúc 16:00.",
        fill=PALE_GOLD,
        border="A16B00",
    )
    add_heading(doc, "10.3. Chuẩn bị để demo không phụ thuộc mạng", 2)
    for item in [
        "Có chế độ MockProvider cho Calendar, SMS và Inventory; cùng interface với adapter thật.",
        "Seed sẵn 3 sale, 6 căn, 2 xe và 1 xung đột để luồng luôn tái hiện được.",
        "Có nút “Inject timeout” và “Create conflict” chỉ xuất hiện ở môi trường demo.",
        "Lưu ảnh/video dự phòng của luồng chuẩn trong trường hợp mạng hoặc API key có vấn đề.",
        "Reset demo bằng migration/seed idempotent, không xóa dữ liệu thủ công khi đang trình bày.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "10.4. Tiêu chí chấm demo", 2)
    demo_checks = [
        ["Hội thoại", "Hiểu yêu cầu và hỏi lại phần thiếu, không bịa lịch/căn."],
        ["Tool use", "Hiển thị được tool/state hoặc log minh họa rõ ràng."],
        ["HITL", "Sale duyệt/sửa/từ chối và luồng tiếp tục đúng."],
        ["Concurrency", "Xung đột không tạo hai booking/hold."],
        ["Reliability", "Lỗi API có trạng thái, retry và thông điệp rõ."],
        ["UX", "Khách biết đang chờ ai, hold còn bao lâu và thao tác tiếp theo."],
    ]
    add_table(doc, ["Tiêu chí", "Dấu hiệu đạt"], demo_checks, [2100, 7260])

    add_heading(doc, "11. Đánh giá, phân công và lộ trình", 1)
    add_heading(doc, "11.1. Bộ chỉ số đánh giá", 2)
    metrics = [
        ["Booking success rate", "Số lịch CONFIRMED / số yêu cầu đủ điều kiện", "≥ 70% trên bộ test"],
        ["Conflict prevention", "Số double-booking lọt qua", "0"],
        ["HITL approval latency", "Thời gian từ pending đến quyết định", "P50 < 2 phút (demo)"],
        ["Calendar sync success", "Đồng bộ thành công không cần can thiệp", "≥ 99% sau retry"],
        ["Hold leak rate", "Hold quá hạn nhưng chưa release", "0"],
        ["Tool-call correctness", "Tool và đối số đúng theo ground truth", "≥ 95%"],
        ["No-show rate", "Khách không đến / lịch đã xác nhận", "Theo dõi xu hướng"],
        ["End-to-end latency", "Yêu cầu → options; approval → confirmed", "P95 theo SLA nhóm đặt"],
    ]
    add_table(doc, ["Chỉ số", "Cách tính", "Mục tiêu ban đầu"], metrics, [2300, 4660, 2400])

    add_heading(doc, "11.2. Bộ test nên có", 2)
    for item in [
        "Câu rõ ràng, câu thiếu ngày, câu mơ hồ “cuối tuần”, và câu đổi ý nhiều lần.",
        "Không có sale rảnh, không có căn phù hợp, chỉ có phương án gần đúng.",
        "Hai khách đồng thời chọn cùng căn/khung giờ.",
        "Sale sửa giờ sau khi khách chọn; approval cũ không được tái sử dụng.",
        "Calendar timeout trước/sau khi tạo event; retry không được sinh event trùng.",
        "Hold hết hạn đúng lúc khách yêu cầu dời; trạng thái cuối phải nhất quán.",
        "Prompt injection nằm trong mô tả căn hoặc tin nhắn khách; Agent không vượt allowlist.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "11.3. Phân công nhóm gợi ý", 2)
    team_rows = [
        ["Frontend", "Chat, options, sale approval card, dashboard, WebSocket."],
        ["Backend", "Auth, booking state machine, transaction, outbox, adapters."],
        ["AI/Agent", "Prompt, tools schema, LangGraph/HITL, eval hội thoại."],
        ["Data/DevOps", "Schema, seed, worker, Docker/Fly.io, logging/metrics."],
        ["QA/Tài liệu", "Test concurrency/lỗi, kịch bản demo, báo cáo và slide."],
    ]
    add_table(doc, ["Nhóm việc", "Đầu ra"], team_rows, [2100, 7260])

    add_heading(doc, "11.4. Lộ trình ba pha", 2)
    roadmap = [
        ["Pha 1 - Demo ổn định", "Mock adapters, hội thoại, HITL, transaction/hold, WebSocket", "Chứng minh end-to-end"],
        ["Pha 2 - Tích hợp thật", "Google/Outlook OAuth, email/SMS, outbox retry, audit", "Pilot nội bộ"],
        ["Pha 3 - Nâng cao", "Route optimization, memory consent, map, eval/no-show", "Tối ưu vận hành"],
    ]
    add_table(doc, ["Pha", "Phạm vi", "Mục tiêu"], roadmap, [2050, 5150, 2160])

    doc.add_page_break()
    add_heading(doc, "12. Kết luận và checklist bàn giao", 1)
    add_body(
        doc,
        "Đề tài khả thi và có giá trị trình diễn cao nếu nhóm giữ đúng trọng tâm: hội thoại tự "
        "nhiên chỉ là lớp vào; điểm khó và đáng chấm nằm ở điều phối tài nguyên, HITL, chống "
        "double-booking, nhất quán khi API lỗi và khả năng đo kết quả.",
    )
    add_callout(
        doc,
        "Thông điệp thuyết trình",
        "“AI không tự ý chốt lịch. AI thu thập nhu cầu và đề xuất; sale phê duyệt; backend khóa "
        "giao dịch, giữ căn có hạn và đồng bộ nhiều hệ thống một cách có kiểm soát.”",
        fill=PALE_GREEN,
        border="2F855A",
    )
    add_heading(doc, "Checklist trước khi gửi/đi demo", 2)
    checklist = [
        "Luồng chuẩn chạy được từ khách → sale → hold → Calendar → xác nhận.",
        "Có dữ liệu seed và nút reset demo.",
        "Có một kịch bản conflict và một kịch bản API timeout.",
        "Approval cũ/stale bị từ chối; retry không tạo bản ghi trùng.",
        "UI hiển thị rõ trạng thái và thời hạn soft-hold.",
        "Không để API key, token hay dữ liệu khách thật trong repo/log.",
        "Có số liệu eval tối thiểu và ảnh/video dự phòng.",
        "Thống nhất thuật ngữ: yêu cầu, phê duyệt, soft-hold, xác nhận, hết hạn, đồng bộ lỗi.",
    ]
    for item in checklist:
        add_bullet(doc, "□ " + item, bullet_id)

    add_heading(doc, "Nguồn tham khảo", 1)
    add_body(
        doc,
        "Các nguồn dưới đây là tài liệu chính thức của sản phẩm/nền tảng, truy cập ngày "
        "28/07/2026. Nội dung trong báo cáo được diễn giải và chuyển thành khuyến nghị thiết kế "
        "cho bối cảnh bài tập lớn.",
    )
    for idx, (publisher, title, url) in enumerate(SOURCES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        set_paragraph_spacing(p, after=7, line=1.08)
        lead = p.add_run(f"[{idx}] {publisher}. ")
        set_run_font(lead, size=10, bold=True, color=INK)
        title_run = p.add_run(f"{title}. ")
        set_run_font(title_run, size=10, color=BLACK, italic=True)
        add_hyperlink(p, "Mở tài liệu", url)

    add_heading(doc, "Ghi chú sử dụng tài liệu", 2)
    add_body(
        doc,
        "Nhóm có thể sao chép các phần Tóm tắt điều hành, Kiến trúc, Kịch bản demo và Nguồn "
        "tham khảo vào báo cáo chung. Khi triển khai thật, cần kiểm tra lại quota, quyền OAuth, "
        "chính sách dữ liệu và phiên bản API tại thời điểm tích hợp.",
    )

    core = doc.core_properties
    core.title = "AI20K - BĐS O2O: Tài liệu tham khảo và đề xuất"
    core.subject = "AI Agent đặt lịch xem nhà, HITL và soft-hold"
    core.author = "Nhóm bài tập lớn AI20K"
    core.keywords = "AI Agent, bất động sản, O2O, HITL, soft-hold, Calendar, PostgreSQL"
    core.comments = "Tài liệu nghiên cứu phục vụ bài tập lớn"

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
