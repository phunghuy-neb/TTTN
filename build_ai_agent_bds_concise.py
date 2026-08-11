from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "deliverables"
OUT_PATH = OUT_DIR / "Tham_Khao_AI_Agent_BDS_Ban_Ngan_Gon_Chinh_Xac.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "64748B"
BODY = "1F2937"
TABLE_FILL = "E8EEF5"
LIGHT_BLUE = "EFF6FB"
LIGHT_GREEN = "EDF7F1"
LIGHT_GOLD = "FFF8E6"
BORDER = "CBD5E1"

SOURCES = [
    (
        "EliseAI - Getting Started with LeasingAI",
        "https://support.meetelise.com/hc/en-us/articles/42925263474061-Getting-Started-with-LeasingAI",
    ),
    (
        "OJO Labs - Conversational AI kết nối dữ liệu bất động sản",
        "https://ojo.com/article/news-room/1WKdlZlj0CF2xUig9LnI1J",
    ),
    (
        "ShowingTime+ - Essential Guide to Real Estate Showings",
        "https://showingtimeplus.com/hubfs/21555945/Gated%20Content/Guides/"
        "ShowingTimePlus_EssentialGuideToRealEstateShowingsForNewAgents.pdf",
    ),
    (
        "Google Calendar API - FreeBusy",
        "https://developers.google.com/workspace/calendar/api/v3/reference/freebusy/query",
    ),
    (
        "Google Calendar API - Create events",
        "https://developers.google.com/workspace/calendar/api/guides/create-events",
    ),
    (
        "OpenAI - Function calling",
        "https://developers.openai.com/api/docs/guides/function-calling",
    ),
    (
        "LangGraph - Interrupts (Human-in-the-loop)",
        "https://docs.langchain.com/oss/python/langgraph/interrupts",
    ),
    (
        "PostgreSQL - Explicit locking",
        "https://www.postgresql.org/docs/current/explicit-locking.html",
    ),
]


def set_run_font(run, size=11, color=BODY, bold=None, italic=None):
    run.font.name = "Calibri"
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Calibri")
    r_pr.rFonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)


def create_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(x.get(qn("w:abstractNumId")))
        for x in numbering.findall(qn("w:abstractNum"))
        if x.get(qn("w:abstractNumId"))
    ]
    num_ids = [
        int(x.get(qn("w:numId")))
        for x in numbering.findall(qn("w:num"))
        if x.get(qn("w:numId"))
    ]
    next_abs = max(abstract_ids or [0]) + 1
    next_num = max(num_ids or [0]) + 1

    def add_one(kind, abstract_id, num_id):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        lvl.append(text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)

        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        lvl.append(p_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        numbering.append(num)

    add_one("bullet", next_abs, next_num)
    add_one("decimal", next_abs + 1, next_num + 1)
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def add_body(doc, text, refs=None):
    p = doc.add_paragraph(style="Normal")
    set_spacing(p)
    r = p.add_run(text)
    set_run_font(r)
    if refs:
        rr = p.add_run(" " + "".join(f"[{x}]" for x in refs))
        set_run_font(rr, color=BLUE, bold=True)
    return p


def add_bullet(doc, text, bullet_id, refs=None):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_spacing(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r)
    if refs:
        rr = p.add_run(" " + "".join(f"[{x}]" for x in refs))
        set_run_font(rr, color=BLUE, bold=True)
    return p


def add_number(doc, text, number_id):
    p = doc.add_paragraph()
    apply_numbering(p, number_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    set_spacing(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.06)
    p.paragraph_format.right_indent = Inches(0.04)
    set_spacing(p, before=4, after=8, line=1.20)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=INK, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    return p


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_split(row)
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)
        borders.append(node)
    tbl_pr.append(borders)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, after=0, line=1.0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=INK, bold=True)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            p = row.cells[i].paragraphs[0]
            set_spacing(p, after=0, line=1.08)
            r = p.add_run(text)
            set_run_font(r, size=9.4)
    spacer = doc.add_paragraph()
    set_spacing(spacer, after=1)
    return table


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Calibri")
    r_pr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    set_spacing(header, after=0, line=1.0)
    hr = header.add_run("AI20K  |  BĐS O2O - AI Agent đặt lịch xem nhà")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_spacing(footer, after=0, line=1.0)
    fr = footer.add_run("Bản tham khảo ngắn gọn  •  Trang ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_page_field(footer)


def add_title_block(doc):
    spacer = doc.add_paragraph()
    set_spacing(spacer, after=10)

    kicker = doc.add_paragraph()
    set_spacing(kicker, after=3, line=1.0)
    r = kicker.add_run("TÀI LIỆU THAM KHẢO NGẮN GỌN")
    set_run_font(r, size=10.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    set_spacing(title, after=5, line=1.0)
    r = title.add_run("AI AGENT ĐẶT LỊCH XEM BẤT ĐỘNG SẢN")
    set_run_font(r, size=23, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    set_spacing(subtitle, after=12, line=1.05)
    r = subtitle.add_run("Tham khảo cách làm bên ngoài và đề xuất cho demo của nhóm")
    set_run_font(r, size=13, color=MUTED, italic=True)

    for label, value in [
        ("Đề tài", "BĐS - Kinh doanh O2O (Doanh nghiệp bất động sản X)"),
        ("Mục đích", "Gửi thành viên nhóm tổng hợp vào báo cáo và dùng khi trình bày demo"),
        ("Cập nhật", "28/07/2026"),
    ]:
        p = doc.add_paragraph()
        set_spacing(p, after=2, line=1.05)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5)


def build():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)
    bullet_id, number_id = create_numbering(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "Ý chính",
        "Đây không chỉ là chatbot tư vấn nhà. Agent phải hiểu yêu cầu của khách, kiểm tra dữ "
        "liệu thật, chờ sale duyệt, rồi backend mới giữ căn và tạo lịch.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )

    add_heading(doc, "1. Các hệ thống bên ngoài đang làm thế nào?", 1)
    add_body(
        doc,
        "Các sản phẩm trên thị trường không công khai toàn bộ kiến trúc bên trong. Vì vậy, "
        "nhóm chỉ nên rút ra cách họ tổ chức trải nghiệm và dữ liệu, không khẳng định họ dùng "
        "đúng một công nghệ cụ thể nào.",
    )
    external_rows = [
        [
            "EliseAI",
            "Trao đổi với khách thuê qua chat, SMS, email/voice; hỗ trợ đặt tour và kết nối CRM/PMS.",
            "Agent phải nối được với dữ liệu vận hành, không chỉ trả lời bằng kiến thức của LLM. [1]",
        ],
        [
            "OJO",
            "Hiểu nhu cầu tìm nhà bằng hội thoại và lấy thông tin từ kho dữ liệu bất động sản.",
            "Tách phần “hiểu câu khách” khỏi phần “lấy dữ liệu nhà”. [2]",
        ],
        [
            "ShowingTime+",
            "Quản lý yêu cầu xem nhà, bước chấp thuận và trao đổi qua nhiều kênh.",
            "Yêu cầu đặt lịch và lịch đã xác nhận phải là hai trạng thái khác nhau. [3]",
        ],
        [
            "Google Calendar",
            "Kiểm tra bận/rảnh, tạo sự kiện, thêm khách mời, địa điểm và nhắc lịch.",
            "Dùng Calendar làm kênh đồng bộ; booking nội bộ vẫn là nguồn dữ liệu chính. [4][5]",
        ],
    ]
    add_table(
        doc,
        ["Hệ thống", "Họ làm gì", "Nhóm học được gì"],
        external_rows,
        [1550, 3600, 4210],
    )

    add_heading(doc, "2. Luồng đặt lịch nên làm", 1)
    for step in [
        "Khách nói tự nhiên, ví dụ: “Chiều thứ Bảy tôi muốn xem căn 2 phòng ngủ ở khu Đông, có xe đón ở quận 1.”",
        "Agent hỏi thêm phần còn thiếu, rồi gọi tool để tìm căn, xem lịch sale, xe và phòng chờ.",
        "Hệ thống đưa ra 2-3 phương án thực sự còn khả dụng để khách chọn.",
        "Khách chọn một phương án. Lúc này chỉ tạo yêu cầu PENDING_SALE_APPROVAL, chưa giữ căn.",
        "Sale nhận thẻ duyệt và chọn Duyệt, Sửa hoặc Từ chối.",
        "Sau khi sale duyệt, backend kiểm tra lại một lần nữa. Nếu vẫn hợp lệ, backend ghi cuộc hẹn và soft-hold trong một transaction ngắn.",
        "Hệ thống đồng bộ sang Google/Outlook Calendar, gửi xác nhận và nhắc lịch.",
        "Nếu khách dời/hủy hoặc hold hết hạn, hệ thống cập nhật trạng thái và giải phóng tài nguyên.",
    ]:
        add_number(doc, step, number_id)

    add_callout(
        doc,
        "Điểm phải làm đúng",
        "Không giữ căn trước khi sale duyệt. Trong lúc chờ sale, căn và lịch có thể thay đổi, "
        "nên backend bắt buộc kiểm tra lại trước khi ghi dữ liệu.",
        fill=LIGHT_GREEN,
        border="2F855A",
    )

    add_heading(doc, "3. Dữ liệu nhà/căn lấy ở đâu?", 1)
    add_body(
        doc,
        "Với bài tập lớn, cách an toàn và ổn định nhất là tự tạo khoảng 30 căn mẫu trong "
        "PostgreSQL. Không cần crawl hàng nghìn tin đăng thật. Điều quan trọng là dữ liệu có "
        "trạng thái rõ để demo được xung đột.",
    )
    for item in [
        "Thông tin căn: mã căn, dự án, khu vực, loại căn, số phòng, diện tích, giá, địa chỉ, ảnh và tiện ích.",
        "Trạng thái: AVAILABLE, SOFT_HOLD, BOOKED, UNAVAILABLE.",
        "Thông tin phục vụ xem nhà: thời lượng xem, khung giờ có thể xem, sale phụ trách và phiên bản dữ liệu.",
        "Bảng liên quan: sales, appointments, approvals, soft_holds, vehicles và audit_logs.",
        "Dữ liệu thật sau này có thể lấy từ CRM, file Excel/Google Sheets hoặc API inventory của doanh nghiệp X.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_body(
        doc,
        "Agent không đọc database trực tiếp. Agent chỉ gọi các tool như search_units, "
        "check_unit_availability và check_sale_availability; backend mới là nơi kiểm tra quyền "
        "và đọc dữ liệu.",
        refs=[6],
    )

    add_heading(doc, "4. Soft-hold và chống trùng lịch", 1)
    add_body(
        doc,
        "Soft-hold là bản ghi giữ căn tạm thời, ví dụ 15 phút. Nó phải có expires_at và status "
        "để hệ thống biết khi nào cần giải phóng. Row-lock của PostgreSQL chỉ dùng trong lúc "
        "transaction đang chạy; không được giữ khóa 15 phút để chờ khách hoặc sale.",
        refs=[8],
    )
    add_heading(doc, "Cách xử lý ngắn gọn", 2)
    for step in [
        "Bắt đầu transaction và khóa hàng căn cần giữ bằng SELECT ... FOR UPDATE NOWAIT.",
        "Kiểm tra lại căn, lịch sale/xe/phòng chờ và các hold còn hiệu lực.",
        "Nếu có xung đột, rollback và trả phương án khác.",
        "Nếu hợp lệ, ghi appointment, soft_hold và commit ngay.",
        "Một job định kỳ đổi hold quá hạn sang EXPIRED; job này không “mở row-lock”.",
    ]:
        add_number(doc, step, number_id)
    add_body(
        doc,
        "Với lịch sale, xe và phòng chờ, nên có thêm ràng buộc ở database để chặn hai cuộc "
        "hẹn bị chồng thời gian; không nên chỉ kiểm tra bằng code.",
    )

    add_heading(doc, "5. Agent, bước sale duyệt (HITL) và đồng bộ lịch", 1)
    for item in [
        "LLM chỉ hiểu câu khách và đề xuất tool call. Backend mới thực thi, validate và kiểm tra quyền. [6]",
        "LangGraph interrupt dùng để dừng luồng ở bước chờ sale. Tác vụ tạo lịch/giữ căn nên đặt sau interrupt để tránh chạy lặp khi resume. [7]",
        "Google Calendar có FreeBusy để xem lịch trống và Events API để tạo lịch. Có thể dùng OAuth hoặc calendar được chia sẻ; Service Account cần được cấp quyền phù hợp. [4][5]",
        "Khi retry tạo sự kiện, phải dùng cùng idempotency key/event ID để tránh tạo hai lịch giống nhau. [5]",
        "Không gọi Calendar bên trong transaction giữ căn. Hãy commit dữ liệu nội bộ trước, rồi dùng worker/hàng đợi để đồng bộ.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "6. Tech stack đề xuất", 1)
    stack_rows = [
        ["Next.js", "Chat cho khách, thẻ sale duyệt, trang điều phối và hiển thị trạng thái."],
        ["FastAPI", "API nghiệp vụ, đăng nhập/phân quyền, tools, WebSocket và xử lý transaction."],
        ["LangGraph", "Giữ trạng thái hội thoại và tạm dừng ở bước HITL."],
        ["PostgreSQL", "Căn, lịch, approval, soft-hold, trạng thái hội thoại và dữ liệu audit."],
        ["Google/Outlook Calendar", "Kiểm tra bận/rảnh và đồng bộ sự kiện."],
        ["Worker + hàng đợi", "Retry Calendar/thông báo và giải phóng hold hết hạn."],
        ["Fly.io / Docker", "Dùng để deploy; demo local bằng Docker Compose cũng được."],
    ]
    add_table(doc, ["Thành phần", "Dùng để làm gì"], stack_rows, [2300, 7060])

    add_heading(doc, "7. Lỗi và bảo mật cần xử lý", 1)
    for item in [
        "Calendar hoặc inventory timeout: chờ tối đa vài giây, retry có giới hạn, sau đó chuyển sang chờ điều phối viên.",
        "Căn vừa bị người khác giữ: không báo thành công; trả lựa chọn gần nhất cho khách.",
        "Sale không duyệt kịp: yêu cầu hết hạn, không tạo hold.",
        "Không gửi toàn bộ số điện thoại/email vào prompt nếu không cần; dùng customer_id hoặc contact_ref.",
        "Sale chỉ được xem khách và lịch mình phụ trách; mọi lần duyệt, gia hạn, dời/hủy đều có audit.",
        "Giao diện phải nói rõ “đang chờ sale”, “đang đồng bộ” hay “đã xác nhận”, tránh làm khách hiểu nhầm.",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "8. Kịch bản demo gọn (5-7 phút)", 1)
    for step in [
        "Khách đăng nhập và nhập yêu cầu xem căn bằng một câu tự nhiên.",
        "Agent gọi tool và trả ba lựa chọn.",
        "Khách chọn; màn hình chuyển sang “Chờ sale xác nhận”.",
        "Sale đăng nhập, sửa hoặc duyệt lịch.",
        "Backend re-check, tạo soft-hold và đồng bộ Calendar; khách nhận cập nhật real-time.",
        "Tạo sẵn một xung đột để chứng minh hệ thống không double-booking mà tự đề xuất lại.",
        "Nếu còn thời gian, demo dời/hủy hoặc lỗi Calendar.",
    ]:
        add_number(doc, step, number_id)

    add_callout(
        doc,
        "Phạm vi nên ưu tiên",
        "Làm chắc luồng cơ bản trước. Tối ưu lộ trình, bản đồ tiện ích, memory sở thích và đo "
        "no-show để ở phần nâng cao; không cần cố nhét hết vào demo đầu.",
        fill=LIGHT_GOLD,
        border="A16B00",
    )

    add_heading(doc, "Nguồn tham khảo", 1)
    intro = doc.add_paragraph()
    set_spacing(intro, before=4, after=4, line=1.0)
    r = intro.add_run("Truy cập ngày 28/07/2026. Bấm vào “Mở nguồn” để xem tài liệu gốc.")
    set_run_font(r, size=9.5, color=MUTED, italic=True)
    for i, (title, url) in enumerate(SOURCES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.30)
        p.paragraph_format.first_line_indent = Inches(-0.30)
        set_spacing(p, after=5, line=1.08)
        r1 = p.add_run(f"[{i}] {title}. ")
        set_run_font(r1, size=10, color=BODY)
        add_hyperlink(p, "Mở nguồn", url)

    core = doc.core_properties
    core.title = "Tham khảo AI Agent đặt lịch xem bất động sản - bản ngắn gọn"
    core.subject = "BĐS O2O, HITL, soft-hold và Calendar"
    core.author = "Nhóm bài tập lớn AI20K"
    core.comments = "Bản viết ngắn gọn, ngôn ngữ tự nhiên"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
